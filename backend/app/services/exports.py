"""Выгрузка данных клуба в Excel (расширение по просьбе клуба).

Аналитика в приложении отвечает на вопросы, которые мы придумали заранее.
Клубу регулярно нужны и другие — «а кто ходил чаще всех в мае», «а какие
жанры любят те, кто ни разу не пришёл». Заранее их не предугадать, поэтому
рядом с готовыми графиками есть сырые таблицы: один файл на тему, а кто
посчитает по ним сводную — решает уже человек.

Устройство: реестр наборов (`DATASETS`), каждый набор — несколько листов,
каждый лист — заголовки и строки. Никакой разметки в запросах: SQL отдаёт
значения, а формат, ширина колонок и часовой пояс достаются им на общем
проходе в `workbook()`. Добавить таблицу значит дописать одну функцию.

Даты пишутся в зоне клуба и БЕЗ смещения: Excel не умеет timezone-aware
и на попытку записать такую дату падает, а человек, открывший файл, всё равно
читает «19:00», а не «16:00 UTC».
"""

import io
import re
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any
from zoneinfo import ZoneInfo

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased

from app.models import (
    Achievement,
    Attendance,
    Availability,
    Confirmation,
    Favourite,
    Feedback,
    Film,
    FilmRating,
    FilmRequest,
    FilmSkip,
    FilmVote,
    Friendship,
    Hall,
    Interest,
    InviteCode,
    Referral,
    Round,
    Screening,
    ShortlistItem,
    Slot,
    Tournament,
    TournamentMatch,
    TournamentOption,
    TournamentVote,
    User,
    Watch,
)
from app.models.enums import ConfirmationState, InterestKind, ScreeningStatus
from app.services import achievements as achievements_service
from app.services import tournaments as tournaments_service

# --- Человеческие подписи вместо значений enum -------------------------------
#
# В файле, который откроет не программист, «waitlist» и «auto_coverage» —
# это не данные, а загадка.

ROLE_LABEL = {
    "user": "участник",
    "moderator": "модератор",
    "admin": "админ",
    "superadmin": "главный админ",
}

INTEREST_LABEL = {"wishlist": "Желаемое", "soon": "Ближайшее"}

REVOKE_LABEL = {
    "manual": "снял сам",
    "expired": "истекла",
    "watched": "посмотрел",
    "superseded": "заменена другой",
}

SCREENING_STATUS_LABEL = {
    "scheduled": "назначен",
    "cancelled": "отменён",
    "completed": "прошёл",
}

CONFIRMATION_LABEL = {
    "confirmed": "придёт",
    "waitlist": "лист ожидания",
    "cancelled": "отменил",
}

ATTENDANCE_LABEL = {"qr": "QR-код", "code": "код с экрана", "manual": "отметил модератор"}

ROUND_STAGE_LABEL = {
    "collecting": "собираем интерес",
    "shortlist_review": "шорт-лист собран",
    "slot_voting": "идёт голосование",
    "schedule_review": "расставляем показы",
    "published": "расписание опубликовано",
    "running": "неделя показов",
    "closed": "закрыт",
}

SHORTLIST_SOURCE_LABEL = {
    "auto_weight": "автопилот, по весу",
    "auto_coverage": "автопилот, по охвату",
    "admin": "поставил админ",
}

TOURNAMENT_STATUS_LABEL = {
    "draft": "черновик",
    "running": "идёт",
    "finished": "завершён",
    "cancelled": "отменён",
}

FILM_STATUS_LABEL = {
    "active": "в каталоге",
    "hidden": "скрыт",
    "pending_moderation": "на модерации",
}

FILM_REQUEST_LABEL = {"pending": "ждёт разбора", "approved": "принята", "rejected": "отклонена"}

WATCH_SOURCE_LABEL = {"manual": "отметил сам", "attendance": "был на показе"}

RATING_SOURCE_LABEL = {"catalog": "из каталога", "screening": "после показа"}


def label(mapping: dict[str, str], value: Any) -> str | None:
    """Подпись по значению enum. Незнакомое отдаём как есть: новый статус
    должен попасть в файл, а не исчезнуть из него."""
    if value is None:
        return None
    return mapping.get(str(value), str(value))


def stars(half_points: int | None) -> float | None:
    """Оценки хранятся полубаллами (1..10); в таблице привычнее звёзды."""
    return None if half_points is None else half_points / 2


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    headers: tuple[str, ...]
    rows: list[tuple]


@dataclass(frozen=True, slots=True)
class Dataset:
    key: str
    # Подпись кнопки в боте. Короткая намеренно: кнопки стоят по две в ряд,
    # и длинное название Telegram обрезает многоточием.
    title: str
    # Что внутри — строкой в меню, рядом с кнопкой.
    summary: str
    # Имя файла без расширения.
    filename: str
    build: Callable[[AsyncSession], Awaitable[list[Sheet]]]


# --- Сборка книги ------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")

DATETIME_FORMAT = "DD.MM.YYYY HH:MM"
DATE_FORMAT = "DD.MM.YYYY"

# Excel обрежет сам, но молча и с потерей смысла — лучше обрезать осознанно.
MAX_SHEET_NAME = 31
MAX_COLUMN_WIDTH = 55
# Ячейка Excel держит 32767 символов; отзыв длиннее — не отзыв, а авария.
MAX_CELL_TEXT = 32_000


def sheet_name(raw: str, taken: set[str]) -> str:
    """Имя листа: без запрещённых символов, не длиннее 31, без повторов."""
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", raw).strip() or "Лист"
    cleaned = cleaned[:MAX_SHEET_NAME]
    name, suffix = cleaned, 2
    while name.lower() in taken:
        tail = f" {suffix}"
        name = cleaned[: MAX_SHEET_NAME - len(tail)] + tail
        suffix += 1
    taken.add(name.lower())
    return name


def _cell_value(value: Any, tz: ZoneInfo) -> Any:
    if isinstance(value, datetime):
        # Наивные даты уже локальные (их некуда переводить), aware — приводим
        # к зоне клуба и срезаем смещение: Excel с ним не работает.
        return value.astimezone(tz).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, str) and len(value) > MAX_CELL_TEXT:
        return value[:MAX_CELL_TEXT] + "…"
    return value


def workbook(sheets: Sequence[Sheet], tz: ZoneInfo) -> bytes:
    """Книга целиком в памяти: файлы клуба измеряются мегабайтами, и временный
    файл на диске ради них — лишняя сущность и лишняя уборка."""
    book = Workbook()
    book.remove(book.active)
    taken: set[str] = set()

    for sheet in sheets:
        page = book.create_sheet(sheet_name(sheet.name, taken))
        page.append(list(sheet.headers))
        widths = [len(header) + 2 for header in sheet.headers]

        for row in sheet.rows:
            page.append([_cell_value(value, tz) for value in row])
            for index, value in enumerate(row[: len(widths)]):
                widths[index] = max(widths[index], min(len(str(value or "")) + 2, MAX_COLUMN_WIDTH))

        for cell in page[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")

        for index, width in enumerate(widths, start=1):
            page.column_dimensions[get_column_letter(index)].width = width

        for column in page.iter_cols(min_row=2):
            for cell in column:
                if isinstance(cell.value, datetime):
                    cell.number_format = DATETIME_FORMAT
                elif isinstance(cell.value, date):
                    cell.number_format = DATE_FORMAT

        if sheet.rows:
            # Заголовок остаётся на месте при прокрутке, а фильтры — то, ради
            # чего таблицу и открывают.
            page.freeze_panes = "A2"
            page.auto_filter.ref = page.dimensions

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


# --- Мелкие помощники для запросов -------------------------------------------


def _count(model, *where) -> sa.ScalarSelect:
    return sa.select(sa.func.count()).select_from(model).where(*where).scalar_subquery()


def _who(alias=User) -> tuple:
    """Человек в трёх колонках: id, имя и ник — по любому из них его найдут."""
    return (alias.id, alias.display_name, alias.tg_username)


WHO_HEADERS = ("ID", "Имя", "Ник в Telegram")
FILM_HEADERS = ("Фильм", "Год")


# --- Наборы данных -----------------------------------------------------------


async def build_users(session: AsyncSession) -> list[Sheet]:
    active_interest = Interest.revoked_at.is_(None)
    query = (
        sa.select(
            User.id,
            User.tg_id,
            User.tg_username,
            User.display_name,
            User.role,
            User.is_active,
            User.created_at,
            User.last_seen_at,
            User.onboarded_at,
            User.access_granted_at,
            InviteCode.code,
            _count(
                Interest,
                Interest.user_id == User.id,
                Interest.kind == InterestKind.WISHLIST,
                active_interest,
            ),
            _count(
                Interest,
                Interest.user_id == User.id,
                Interest.kind == InterestKind.SOON,
                active_interest,
            ),
            _count(Watch, Watch.user_id == User.id),
            _count(FilmSkip, FilmSkip.user_id == User.id),
            _count(FilmRating, FilmRating.user_id == User.id),
            sa.select(sa.func.avg(FilmRating.score))
            .where(FilmRating.user_id == User.id)
            .scalar_subquery(),
            _count(Favourite, Favourite.user_id == User.id),
            _count(Friendship, Friendship.user_id == User.id),
            _count(Friendship, Friendship.friend_id == User.id),
            _count(
                Confirmation,
                Confirmation.user_id == User.id,
                Confirmation.state == ConfirmationState.CONFIRMED,
            ),
            _count(Attendance, Attendance.user_id == User.id),
            _count(Feedback, Feedback.user_id == User.id),
            _count(Achievement, Achievement.user_id == User.id),
            _count(FilmVote, FilmVote.user_id == User.id),
            _count(Availability, Availability.user_id == User.id),
            _count(TournamentVote, TournamentVote.user_id == User.id),
            _count(Referral, Referral.referrer_id == User.id),
        )
        .select_from(User)
        .outerjoin(InviteCode, InviteCode.id == User.invite_code_id)
        .order_by(User.id)
    )

    rows = []
    for row in (await session.execute(query)).all():
        (
            user_id, tg_id, username, name, role, is_active, created, seen, onboarded,
            access, code, wishlist, soon, watched, skipped, ratings, avg_score, favourites,
            following, followers, confirmed, attended, feedback, awards, film_votes,
            evenings, tournament_votes, invited,
        ) = row
        rows.append(
            (
                user_id, tg_id, username, name, label(ROLE_LABEL, role),
                "да" if is_active else "нет", created, seen, onboarded, access, code,
                wishlist, soon, watched, skipped, ratings,
                round(float(avg_score) / 2, 2) if avg_score is not None else None,
                favourites, following, followers, confirmed, attended, feedback, awards,
                film_votes, evenings, tournament_votes, invited,
            )
        )

    return [
        Sheet(
            "Пользователи",
            (
                "ID", "Telegram ID", "Ник", "Имя", "Роль", "Активен", "Зарегистрирован",
                "Последний визит", "Знакомство пройдено", "Доступ выдан", "Код приглашения",
                "«Желаемое»", "«Ближайшее»", "Просмотрено", "Пролистано в ленте", "Оценок",
                "Средняя оценка", "Любимых фильмов", "Подписок", "Подписчиков", "Записей «приду»",
                "Посещений", "Отзывов", "Достижений", "Голосов за фильмы", "Отмечено вечеров",
                "Голосов в турнирах", "Позвал друзей",
            ),
            rows,
        )
    ]


async def build_ratings(session: AsyncSession) -> list[Sheet]:
    query = (
        sa.select(
            *_who(),
            Film.id,
            Film.title_ru,
            Film.year,
            FilmRating.score,
            FilmRating.source,
            FilmRating.created_at,
            FilmRating.updated_at,
        )
        .join(User, User.id == FilmRating.user_id)
        .join(Film, Film.id == FilmRating.film_id)
        .order_by(FilmRating.created_at)
    )
    rows = [
        (uid, name, nick, film_id, title, year, stars(score), label(RATING_SOURCE_LABEL, source),
         created, updated)
        for uid, name, nick, film_id, title, year, score, source, created, updated in (
            await session.execute(query)
        ).all()
    ]
    return [
        Sheet(
            "Оценки",
            (*WHO_HEADERS, "ID фильма", *FILM_HEADERS, "Оценка", "Откуда", "Поставлена",
             "Изменена"),
            rows,
        )
    ]


async def build_interests(session: AsyncSession) -> list[Sheet]:
    marks = (
        sa.select(
            *_who(),
            Film.title_ru,
            Film.year,
            Interest.kind,
            Interest.created_at,
            Interest.revoked_at,
            Interest.revoke_reason,
        )
        .join(User, User.id == Interest.user_id)
        .join(Film, Film.id == Interest.film_id)
        .order_by(Interest.created_at)
    )
    mark_rows = [
        (uid, name, nick, title, year, label(INTEREST_LABEL, kind),
         "да" if revoked is None else "нет", created, revoked, label(REVOKE_LABEL, reason))
        for uid, name, nick, title, year, kind, created, revoked, reason in (
            await session.execute(marks)
        ).all()
    ]

    watches = (
        sa.select(*_who(), Film.title_ru, Film.year, Watch.source, Watch.screening_id,
                  Watch.created_at)
        .join(User, User.id == Watch.user_id)
        .join(Film, Film.id == Watch.film_id)
        .order_by(Watch.created_at)
    )
    watch_rows = [
        (uid, name, nick, title, year, label(WATCH_SOURCE_LABEL, source), screening, created)
        for uid, name, nick, title, year, source, screening, created in (
            await session.execute(watches)
        ).all()
    ]

    skips = (
        sa.select(*_who(), Film.title_ru, Film.year, FilmSkip.created_at)
        .join(User, User.id == FilmSkip.user_id)
        .join(Film, Film.id == FilmSkip.film_id)
        .order_by(FilmSkip.created_at)
    )
    skip_rows = [tuple(row) for row in (await session.execute(skips)).all()]

    return [
        Sheet(
            "Отметки",
            (*WHO_HEADERS, *FILM_HEADERS, "Тип", "Активна", "Поставлена", "Снята", "Почему снята"),
            mark_rows,
        ),
        Sheet(
            "Просмотрено",
            (*WHO_HEADERS, *FILM_HEADERS, "Откуда", "ID показа", "Когда"),
            watch_rows,
        ),
        Sheet("Пролистано в ленте", (*WHO_HEADERS, *FILM_HEADERS, "Когда"), skip_rows),
    ]


async def build_films(session: AsyncSession) -> list[Sheet]:
    active = Interest.revoked_at.is_(None)
    query = sa.select(
        Film.id,
        Film.title_ru,
        Film.title_orig,
        Film.year,
        Film.runtime_min,
        Film.genres,
        Film.directors,
        Film.status,
        Film.tmdb_id,
        Film.tmdb_rating,
        Film.tmdb_votes,
        Film.kp_id,
        Film.kp_rating,
        Film.kp_votes,
        Film.kp_top250,
        sa.select(sa.func.avg(FilmRating.score))
        .where(FilmRating.film_id == Film.id)
        .scalar_subquery(),
        _count(FilmRating, FilmRating.film_id == Film.id),
        _count(Interest, Interest.film_id == Film.id, Interest.kind == InterestKind.WISHLIST,
               active),
        _count(Interest, Interest.film_id == Film.id, Interest.kind == InterestKind.SOON, active),
        _count(Watch, Watch.film_id == Film.id),
        _count(FilmSkip, FilmSkip.film_id == Film.id),
        _count(Favourite, Favourite.film_id == Film.id),
        _count(Screening, Screening.film_id == Film.id,
               Screening.status != ScreeningStatus.CANCELLED),
        Film.created_at,
    ).order_by(Film.id)

    rows = []
    for row in (await session.execute(query)).all():
        values = list(row)
        values[7] = label(FILM_STATUS_LABEL, values[7])
        values[15] = round(float(values[15]) / 2, 2) if values[15] is not None else None
        rows.append(tuple(values))

    return [
        Sheet(
            "Каталог",
            (
                "ID", "Название", "Оригинал", "Год", "Минут", "Жанры", "Режиссёры", "Статус",
                "TMDB ID", "Рейтинг TMDB", "Голосов TMDB", "КП ID", "Рейтинг КП", "Голосов КП",
                "Топ-250 КП", "Рейтинг клуба", "Оценок в клубе", "«Желаемое»", "«Ближайшее»",
                "Просмотрели", "Пролистали", "В любимых", "Показов", "Добавлен",
            ),
            rows,
        )
    ]


async def build_cycle(session: AsyncSession) -> list[Sheet]:
    rounds = sa.select(
        Round.id,
        Round.week_start,
        Round.stage,
        Round.shortlist_locked_at,
        Round.schedule_locked_at,
        Round.published_at,
        Round.low_activity,
        Round.skipped_reason,
        _count(ShortlistItem, ShortlistItem.round_id == Round.id),
        _count(Screening, Screening.round_id == Round.id),
        sa.select(sa.func.count(sa.distinct(FilmVote.user_id)))
        .where(FilmVote.round_id == Round.id)
        .scalar_subquery(),
        sa.select(sa.func.count(sa.distinct(Availability.user_id)))
        .where(Availability.round_id == Round.id)
        .scalar_subquery(),
    ).order_by(Round.week_start)
    round_rows = [
        (rid, week, label(ROUND_STAGE_LABEL, stage), locked, scheduled, published,
         "да" if low else "нет", reason, items, screenings, voters, available)
        for rid, week, stage, locked, scheduled, published, low, reason, items, screenings,
        voters, available in (await session.execute(rounds)).all()
    ]

    shortlist = (
        sa.select(
            Round.week_start,
            ShortlistItem.position,
            Film.title_ru,
            Film.year,
            ShortlistItem.source,
            ShortlistItem.weight_snapshot,
            _count(FilmVote, FilmVote.round_id == ShortlistItem.round_id,
                   FilmVote.film_id == ShortlistItem.film_id),
        )
        .join(Round, Round.id == ShortlistItem.round_id)
        .join(Film, Film.id == ShortlistItem.film_id)
        .order_by(Round.week_start, ShortlistItem.position)
    )
    shortlist_rows = [
        (week, position, title, year, label(SHORTLIST_SOURCE_LABEL, source),
         round(weight, 3) if weight is not None else None, votes)
        for week, position, title, year, source, weight, votes in (
            await session.execute(shortlist)
        ).all()
    ]

    votes = (
        sa.select(Round.week_start, *_who(), Film.title_ru, Film.year, FilmVote.created_at)
        .join(Round, Round.id == FilmVote.round_id)
        .join(User, User.id == FilmVote.user_id)
        .join(Film, Film.id == FilmVote.film_id)
        .order_by(Round.week_start, User.id)
    )
    vote_rows = [tuple(row) for row in (await session.execute(votes)).all()]

    evenings = (
        sa.select(Round.week_start, *_who(), Slot.starts_at, Slot.duration_min,
                  Availability.created_at)
        .join(Round, Round.id == Availability.round_id)
        .join(User, User.id == Availability.user_id)
        .join(Slot, Slot.id == Availability.slot_id)
        .order_by(Round.week_start, Slot.starts_at, User.id)
    )
    evening_rows = [tuple(row) for row in (await session.execute(evenings)).all()]

    return [
        Sheet(
            "Циклы",
            ("ID", "Неделя показов", "Стадия", "Шорт-лист заперт", "Расписание заперто",
             "Опубликован", "Мало активности", "Почему пропущен", "В шорт-листе", "Показов",
             "Голосовало за фильмы", "Отметило вечера"),
            round_rows,
        ),
        Sheet(
            "Шорт-листы",
            ("Неделя показов", "Позиция", *FILM_HEADERS, "Откуда", "Вес на срезе", "Голосов"),
            shortlist_rows,
        ),
        Sheet(
            "Голоса за фильмы",
            ("Неделя показов", *WHO_HEADERS, *FILM_HEADERS, "Когда"),
            vote_rows,
        ),
        Sheet(
            "Свободные вечера",
            ("Неделя показов", *WHO_HEADERS, "Вечер", "Минут", "Когда"),
            evening_rows,
        ),
    ]


async def build_screenings(session: AsyncSession) -> list[Sheet]:
    query = (
        sa.select(
            Screening.id,
            Slot.starts_at,
            Slot.duration_min,
            Film.title_ru,
            Film.year,
            Screening.title,
            Round.week_start,
            Screening.is_manual,
            Screening.in_english,
            Screening.status,
            Hall.name,
            Hall.capacity,
            Screening.expected_attendance,
            _count(Confirmation, Confirmation.screening_id == Screening.id,
                   Confirmation.state == ConfirmationState.CONFIRMED),
            _count(Confirmation, Confirmation.screening_id == Screening.id,
                   Confirmation.state == ConfirmationState.WAITLIST),
            _count(Confirmation, Confirmation.screening_id == Screening.id,
                   Confirmation.state == ConfirmationState.CANCELLED),
            _count(Attendance, Attendance.screening_id == Screening.id),
            sa.select(sa.func.avg(Feedback.film_rating))
            .where(Feedback.screening_id == Screening.id)
            .scalar_subquery(),
            _count(Feedback, Feedback.screening_id == Screening.id),
            Screening.registration_url,
            Screening.cancel_reason,
            Screening.note,
            Screening.created_at,
        )
        .join(Slot, Slot.id == Screening.slot_id)
        .join(Hall, Hall.id == Slot.hall_id)
        .outerjoin(Film, Film.id == Screening.film_id)
        .outerjoin(Round, Round.id == Screening.round_id)
        .order_by(Slot.starts_at)
    )

    rows = []
    for row in (await session.execute(query)).all():
        (sid, starts, duration, film, year, title, week, manual, english, status, hall,
         capacity, expected, confirmed, waitlist, cancelled, came, avg_rating, feedback,
         url, reason, note, created) = row
        rows.append(
            (
                sid, starts, duration, film, year, title, week,
                "вручную" if manual else "из цикла", "да" if english else "нет",
                label(SCREENING_STATUS_LABEL, status), hall, capacity, expected, confirmed,
                waitlist, cancelled, came,
                round(came / confirmed * 100) if confirmed else None,
                round(float(avg_rating) / 2, 2) if avg_rating is not None else None,
                feedback, url, reason, note, created,
            )
        )

    confirmations = (
        sa.select(
            Confirmation.screening_id,
            Slot.starts_at,
            Film.title_ru,
            *_who(),
            Confirmation.state,
            Confirmation.created_at,
            Confirmation.cancelled_at,
            Confirmation.was_late_cancel,
        )
        .join(Screening, Screening.id == Confirmation.screening_id)
        .join(Slot, Slot.id == Screening.slot_id)
        .outerjoin(Film, Film.id == Screening.film_id)
        .join(User, User.id == Confirmation.user_id)
        .order_by(Slot.starts_at, User.id)
    )
    confirmation_rows = [
        (sid, starts, film, uid, name, nick, label(CONFIRMATION_LABEL, state), created,
         cancelled, "да" if late else "нет")
        for sid, starts, film, uid, name, nick, state, created, cancelled, late in (
            await session.execute(confirmations)
        ).all()
    ]

    marker = aliased(User)
    attendance = (
        sa.select(
            Attendance.screening_id,
            Slot.starts_at,
            Film.title_ru,
            *_who(),
            Attendance.method,
            Attendance.marked_at,
            marker.display_name,
        )
        .join(Screening, Screening.id == Attendance.screening_id)
        .join(Slot, Slot.id == Screening.slot_id)
        .outerjoin(Film, Film.id == Screening.film_id)
        .join(User, User.id == Attendance.user_id)
        .outerjoin(marker, marker.id == Attendance.marked_by)
        .order_by(Slot.starts_at, User.id)
    )
    attendance_rows = [
        (sid, starts, film, uid, name, nick, label(ATTENDANCE_LABEL, method), marked, by)
        for sid, starts, film, uid, name, nick, method, marked, by in (
            await session.execute(attendance)
        ).all()
    ]

    return [
        Sheet(
            "Показы",
            (
                "ID", "Начало", "Минут", "Фильм", "Год", "Заголовок события", "Неделя цикла",
                "Как назначен", "На английском", "Статус", "Зал", "Вместимость", "Ожидалось",
                "Записалось", "В листе ожидания", "Отменили", "Пришло", "Дошло, %",
                "Оценка показа", "Отзывов", "Ссылка на регистрацию", "Причина отмены",
                "Заметка", "Создан",
            ),
            rows,
        ),
        Sheet(
            "Записи",
            ("ID показа", "Начало", "Фильм", *WHO_HEADERS, "Состояние", "Записался", "Отменил",
             "Поздняя отмена"),
            confirmation_rows,
        ),
        Sheet(
            "Посещения",
            ("ID показа", "Начало", "Фильм", *WHO_HEADERS, "Как отмечен", "Когда", "Кто отметил"),
            attendance_rows,
        ),
    ]


async def build_feedback(session: AsyncSession) -> list[Sheet]:
    query = (
        sa.select(
            Feedback.screening_id,
            Slot.starts_at,
            Film.title_ru,
            *_who(),
            Feedback.film_rating,
            Feedback.review_text,
            Feedback.org_sound,
            Feedback.org_picture,
            Feedback.org_hall,
            Feedback.org_time,
            Feedback.org_comment,
            Feedback.created_at,
        )
        .join(Screening, Screening.id == Feedback.screening_id)
        .join(Slot, Slot.id == Screening.slot_id)
        .outerjoin(Film, Film.id == Feedback.film_id)
        .join(User, User.id == Feedback.user_id)
        .order_by(Slot.starts_at, User.id)
    )
    rows = [
        (sid, starts, film, uid, name, nick, stars(rating), text, sound, picture, hall, time_,
         comment, created)
        for sid, starts, film, uid, name, nick, rating, text, sound, picture, hall, time_,
        comment, created in (await session.execute(query)).all()
    ]
    return [
        Sheet(
            "Отзывы",
            ("ID показа", "Начало", "Фильм", *WHO_HEADERS, "Оценка фильма", "Отзыв", "Звук",
             "Картинка", "Зал", "Время", "Комментарий об организации", "Когда"),
            rows,
        )
    ]


def _round_label(round_no: int, options: int) -> str:
    """«1/8», «полуфинал», «финал» — по числу пар на этапе.

    Пары считаются из размера сетки, а не запросом: сетка всегда степень двойки,
    и на этапе n пар ровно вдвое меньше, чем на предыдущем.
    """
    matches = options // (2**round_no) if options else 0
    return tournaments_service.round_name(matches) if matches >= 1 else f"этап {round_no}"


async def build_tournaments(session: AsyncSession) -> list[Sheet]:
    winner = aliased(TournamentOption)
    author = aliased(User)
    header = (
        sa.select(
            Tournament.id,
            Tournament.title,
            Tournament.description,
            Tournament.status,
            Tournament.current_round,
            Tournament.stage_hours,
            Tournament.started_at,
            Tournament.finished_at,
            winner.title,
            author.display_name,
            _count(TournamentOption, TournamentOption.tournament_id == Tournament.id),
            sa.select(sa.func.count(sa.distinct(TournamentVote.user_id)))
            .select_from(TournamentVote)
            .join(TournamentMatch, TournamentMatch.id == TournamentVote.match_id)
            .where(TournamentMatch.tournament_id == Tournament.id)
            .scalar_subquery(),
            Tournament.created_at,
        )
        .outerjoin(winner, winner.id == Tournament.winner_option_id)
        .outerjoin(author, author.id == Tournament.created_by)
        .order_by(Tournament.id)
    )
    header_rows = [
        (tid, title, description, label(TOURNAMENT_STATUS_LABEL, status), current, hours,
         started, finished, champion, created_by, options, voters, created)
        for tid, title, description, status, current, hours, started, finished, champion,
        created_by, options, voters, created in (await session.execute(header)).all()
    ]

    options = (
        sa.select(
            Tournament.title,
            TournamentOption.seed,
            TournamentOption.title,
            TournamentOption.subtitle,
            Film.title_ru,
            Film.year,
        )
        .join(Tournament, Tournament.id == TournamentOption.tournament_id)
        .outerjoin(Film, Film.id == TournamentOption.film_id)
        .order_by(Tournament.id, TournamentOption.seed)
    )
    option_rows = [tuple(row) for row in (await session.execute(options)).all()]

    side_a = aliased(TournamentOption)
    side_b = aliased(TournamentOption)
    champion = aliased(TournamentOption)
    total = sa.select(sa.func.count(sa.distinct(TournamentOption.id))).where(
        TournamentOption.tournament_id == Tournament.id
    ).scalar_subquery()
    matches = (
        sa.select(
            Tournament.title,
            TournamentMatch.round_no,
            total,
            TournamentMatch.position,
            side_a.title,
            side_b.title,
            _count(TournamentVote, TournamentVote.match_id == TournamentMatch.id,
                   TournamentVote.option_id == TournamentMatch.option_a_id),
            _count(TournamentVote, TournamentVote.match_id == TournamentMatch.id,
                   TournamentVote.option_id == TournamentMatch.option_b_id),
            champion.title,
            TournamentMatch.opens_at,
            TournamentMatch.closes_at,
        )
        .join(Tournament, Tournament.id == TournamentMatch.tournament_id)
        .outerjoin(side_a, side_a.id == TournamentMatch.option_a_id)
        .outerjoin(side_b, side_b.id == TournamentMatch.option_b_id)
        .outerjoin(champion, champion.id == TournamentMatch.winner_option_id)
        .order_by(Tournament.id, TournamentMatch.round_no, TournamentMatch.position)
    )
    match_rows = [
        (name, _round_label(round_no, size or 0), position, a, b,
         votes_a, votes_b, won, opens, closes)
        for name, round_no, size, position, a, b, votes_a, votes_b, won, opens, closes in (
            await session.execute(matches)
        ).all()
    ]

    picked = aliased(TournamentOption)
    votes = (
        sa.select(
            Tournament.title,
            TournamentMatch.round_no,
            TournamentMatch.position,
            *_who(),
            picked.title,
            TournamentVote.created_at,
        )
        .join(TournamentMatch, TournamentMatch.id == TournamentVote.match_id)
        .join(Tournament, Tournament.id == TournamentMatch.tournament_id)
        .join(User, User.id == TournamentVote.user_id)
        .join(picked, picked.id == TournamentVote.option_id)
        .order_by(Tournament.id, TournamentMatch.round_no, TournamentMatch.position, User.id)
    )
    vote_rows = [tuple(row) for row in (await session.execute(votes)).all()]

    return [
        Sheet(
            "Турниры",
            ("ID", "Название", "Описание", "Статус", "Идёт этап", "Часов на этап", "Старт",
             "Финиш", "Победитель", "Создал", "Вариантов", "Голосовало людей", "Создан"),
            header_rows,
        ),
        Sheet("Варианты", ("Турнир", "Посев", "Название", "Подпись", *FILM_HEADERS), option_rows),
        Sheet(
            "Пары",
            ("Турнир", "Этап", "Позиция", "Вариант A", "Вариант B", "Голосов за A",
             "Голосов за B", "Победитель", "Открыт", "Закрыт"),
            match_rows,
        ),
        Sheet(
            "Голоса в турнирах",
            ("Турнир", "Этап №", "Пара", *WHO_HEADERS, "За кого", "Когда"),
            vote_rows,
        ),
    ]


async def build_achievements(session: AsyncSession) -> list[Sheet]:
    granter = aliased(User)
    query = (
        sa.select(
            *_who(),
            Achievement.code,
            Achievement.title,
            Achievement.description,
            Achievement.tier,
            Achievement.earned_at,
            granter.display_name,
        )
        .join(User, User.id == Achievement.user_id)
        .outerjoin(granter, granter.id == Achievement.granted_by)
        .order_by(Achievement.earned_at)
    )

    rows = []
    for uid, name, nick, code, own_title, own_description, own_tier, earned, by in (
        await session.execute(query)
    ).all():
        rule = achievements_service.BY_CODE.get(code)
        tier = own_tier or (rule.tier.value if rule else None)
        rows.append(
            (
                uid, name, nick, code,
                own_title or (rule.title if rule else code),
                own_description or (rule.description if rule else None),
                achievements_service.TIER_LABEL.get(tier, tier),
                "да" if rule and rule.secret else "нет",
                "именная" if by else "по правилу",
                earned, by,
            )
        )

    # Сколько человек дошло до каждой ступени — тот самый вопрос, ради которого
    # ачивки и заводили: если платину не взял никто, планка стоит не там.
    counts = (
        await session.execute(
            sa.select(Achievement.code, sa.func.count())
            .where(Achievement.granted_by.is_(None))
            .group_by(Achievement.code)
        )
    ).all()
    earned_by = dict(counts)
    ladder_rows = [
        (
            achievements_service.GROUP_LABEL.get(rule.group, rule.group),
            achievements_service.TIER_LABEL[rule.tier],
            rule.title,
            rule.description,
            "да" if rule.secret else "нет",
            rule.target,
            earned_by.get(rule.code, 0),
        )
        for rule in achievements_service.RULES
    ]

    return [
        Sheet(
            "Выданные достижения",
            (*WHO_HEADERS, "Код", "Название", "Условие", "Уровень", "Секретная", "Как выдана",
             "Когда", "Кто выдал"),
            rows,
        ),
        Sheet(
            "Реестр достижений",
            ("Цель", "Уровень", "Название", "Условие", "Секретная", "Порог", "Получили"),
            ladder_rows,
        ),
    ]


async def build_social(session: AsyncSession) -> list[Sheet]:
    friend = aliased(User)
    back = aliased(Friendship)
    friendships = (
        sa.select(
            *_who(),
            friend.id,
            friend.display_name,
            friend.tg_username,
            sa.exists().where(
                back.user_id == Friendship.friend_id, back.friend_id == Friendship.user_id
            ),
            Friendship.created_at,
        )
        .join(User, User.id == Friendship.user_id)
        .join(friend, friend.id == Friendship.friend_id)
        .order_by(User.id, friend.id)
    )
    friendship_rows = [
        (uid, name, nick, fid, fname, fnick, "да" if mutual else "нет", created)
        for uid, name, nick, fid, fname, fnick, mutual, created in (
            await session.execute(friendships)
        ).all()
    ]

    favourites = (
        sa.select(*_who(), Favourite.position, Film.title_ru, Film.year, Favourite.created_at)
        .join(User, User.id == Favourite.user_id)
        .join(Film, Film.id == Favourite.film_id)
        .order_by(User.id, Favourite.position)
    )
    favourite_rows = [tuple(row) for row in (await session.execute(favourites)).all()]

    invitee = aliased(User)
    accepted = sa.exists().where(
        Interest.user_id == Referral.invitee_id,
        Interest.film_id == Referral.film_id,
        Interest.revoked_at.is_(None),
    )
    referrals = (
        sa.select(
            *_who(),
            invitee.id,
            invitee.display_name,
            invitee.tg_username,
            Film.title_ru,
            Film.year,
            accepted,
            Referral.created_at,
        )
        .join(User, User.id == Referral.referrer_id)
        .join(invitee, invitee.id == Referral.invitee_id)
        .join(Film, Film.id == Referral.film_id)
        .order_by(Referral.created_at)
    )
    referral_rows = [
        (uid, name, nick, iid, iname, inick, title, year, "да" if marked else "нет", created)
        for uid, name, nick, iid, iname, inick, title, year, marked, created in (
            await session.execute(referrals)
        ).all()
    ]

    return [
        Sheet(
            "Друзья",
            (*WHO_HEADERS, "ID друга", "Друг", "Ник друга", "Взаимно", "Когда"),
            friendship_rows,
        ),
        Sheet(
            "Любимые фильмы",
            (*WHO_HEADERS, "Позиция", *FILM_HEADERS, "Когда"),
            favourite_rows,
        ),
        Sheet(
            "Приглашения на фильм",
            (*WHO_HEADERS, "ID приглашённого", "Приглашённый", "Ник приглашённого", *FILM_HEADERS,
             "Отметил фильм", "Когда"),
            referral_rows,
        ),
    ]


async def build_requests(session: AsyncSession) -> list[Sheet]:
    resolver = aliased(User)
    requests = (
        sa.select(
            FilmRequest.id,
            *_who(),
            FilmRequest.raw_title,
            FilmRequest.raw_year,
            FilmRequest.note,
            FilmRequest.status,
            Film.title_ru,
            resolver.display_name,
            FilmRequest.resolution_comment,
            FilmRequest.created_at,
            FilmRequest.resolved_at,
        )
        .join(User, User.id == FilmRequest.user_id)
        .outerjoin(Film, Film.id == FilmRequest.resolved_film_id)
        .outerjoin(resolver, resolver.id == FilmRequest.resolved_by)
        .order_by(FilmRequest.created_at)
    )
    request_rows = [
        (rid, uid, name, nick, title, year, note, label(FILM_REQUEST_LABEL, status), resolved,
         by, comment, created, resolved_at)
        for rid, uid, name, nick, title, year, note, status, resolved, by, comment, created,
        resolved_at in (await session.execute(requests)).all()
    ]

    author = aliased(User)
    codes = (
        sa.select(
            InviteCode.code,
            author.display_name,
            InviteCode.max_activations,
            _count(User, User.invite_code_id == InviteCode.id),
            InviteCode.note,
            InviteCode.revoked_at,
            InviteCode.created_at,
        )
        .join(author, author.id == InviteCode.created_by)
        .order_by(InviteCode.created_at)
    )
    code_rows = [tuple(row) for row in (await session.execute(codes)).all()]

    return [
        Sheet(
            "Заявки на фильмы",
            ("ID", *WHO_HEADERS, "Что просили", "Год", "Комментарий", "Статус", "Какой фильм",
             "Кто разобрал", "Решение", "Создана", "Разобрана"),
            request_rows,
        ),
        Sheet(
            "Коды приглашений",
            ("Код", "Создал", "Активаций всего", "Использовано", "Заметка", "Отозван", "Создан"),
            code_rows,
        ),
    ]


DATASETS: tuple[Dataset, ...] = (
    Dataset(
        "users", "👥 Люди", "все участники и что каждый успел сделать",
        "Пользователи", build_users,
    ),
    Dataset(
        "ratings", "⭐ Оценки", "кто какому фильму сколько поставил",
        "Оценки", build_ratings,
    ),
    Dataset(
        "interests", "🔖 Отметки", "«желаемое», «ближайшее», просмотренное, пролистанное",
        "Отметки", build_interests,
    ),
    Dataset(
        "films", "🎞 Каталог", "фильмы с рейтингами и спросом на них",
        "Каталог", build_films,
    ),
    Dataset(
        "cycle", "🗳 Голосование", "циклы, шорт-листы, голоса за фильмы и вечера",
        "Голосование", build_cycle,
    ),
    Dataset(
        "screenings", "📅 Показы", "показы, записи «приду» и кто дошёл",
        "Показы", build_screenings,
    ),
    Dataset(
        "feedback", "💬 Отзывы", "оценки после показа и что писали об организации",
        "Отзывы", build_feedback,
    ),
    Dataset(
        "tournaments", "🏆 Турниры", "сетки, пары и все голоса в них",
        "Турниры", build_tournaments,
    ),
    Dataset(
        "achievements", "🎖 Ачивки", "кто что получил и сколько людей дошло до ступени",
        "Достижения", build_achievements,
    ),
    Dataset(
        "social", "🤝 Друзья", "подписки, любимые фильмы, приглашения по ссылке",
        "Друзья", build_social,
    ),
    Dataset(
        "requests", "📮 Заявки", "заявки на фильмы и коды приглашений",
        "Заявки", build_requests,
    ),
)

BY_KEY: dict[str, Dataset] = {dataset.key: dataset for dataset in DATASETS}

# Ключ кнопки «всё сразу». Не набор: он не строит своих листов, а склеивает
# чужие, и в реестре ему делать нечего.
EVERYTHING = "all"


async def build(session: AsyncSession, key: str) -> list[Sheet]:
    if key == EVERYTHING:
        sheets: list[Sheet] = []
        for dataset in DATASETS:
            sheets.extend(await dataset.build(session))
        return sheets

    dataset = BY_KEY.get(key)
    if dataset is None:
        raise KeyError(key)
    return await dataset.build(session)


def filename(key: str, now: datetime) -> str:
    stem = "Киноклуб — всё" if key == EVERYTHING else f"Киноклуб — {BY_KEY[key].filename}"
    return f"{stem} {now:%Y-%m-%d}.xlsx"


async def export(session: AsyncSession, key: str, tz: ZoneInfo) -> bytes:
    return workbook(await build(session, key), tz)
