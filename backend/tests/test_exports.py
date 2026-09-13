"""Выгрузка данных в Excel и пароль на неё (расширение по просьбе клуба).

Главное, что тут проверяется, — что каждый запрос вообще выполняется и что
число колонок сходится с числом заголовков. Ошибка в одну колонку не падает:
она молча сдвигает данные, и в файле «Оценка» оказывается в графе «Год».
"""

import io
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.services import analytics_gate, exports
from app.services.exports import DATASETS, Sheet, sheet_name, workbook
from tests.conftest import SUPERADMIN_TG_ID, login
from tests.test_analytics import held_screening

MOSCOW = ZoneInfo("Europe/Moscow")


def read(data: bytes):
    return load_workbook(io.BytesIO(data))


async def test_every_dataset_builds_on_an_empty_base(session):
    """Пустой клуб — обычное состояние свежей установки, и файл должен
    получаться с одними заголовками, а не падать."""
    for dataset in DATASETS:
        sheets = await dataset.build(session)
        assert sheets, dataset.key
        assert all(sheet.headers for sheet in sheets), dataset.key


@pytest.mark.parametrize("key", [dataset.key for dataset in DATASETS])
async def test_columns_match_headers(session, key):
    """Строка шире или уже заголовка — это сдвиг данных, а не мелочь."""
    _, _, _, _, _ = await held_screening(session)

    for sheet in await exports.build(session, key):
        for row in sheet.rows:
            assert len(row) == len(sheet.headers), f"{key} / {sheet.name}"


async def test_everything_collects_every_dataset(session):
    await held_screening(session)

    separate = 0
    for dataset in DATASETS:
        separate += len(await dataset.build(session))
    book = read(await exports.export(session, exports.EVERYTHING, MOSCOW))

    assert len(book.sheetnames) == separate
    # Имена листов в Excel уникальны, и склейка не должна их терять.
    assert len(set(book.sheetnames)) == separate


async def test_users_sheet_counts_what_people_actually_did(session):
    _, _, _, _, voters = await held_screening(session)

    book = read(await exports.export(session, "users", MOSCOW))
    page = book["Пользователи"]
    header = [cell.value for cell in page[1]]
    rows = {
        row[0]: dict(zip(header, row, strict=True))
        for row in page.iter_rows(min_row=2, values_only=True)
    }

    came = rows[voters[0].id]
    assert came["Посещений"] == 1
    assert came["Записей «приду»"] == 1
    assert rows[voters[2].id]["Посещений"] == 0


async def test_dates_are_written_in_club_time_without_offset(session):
    """Excel не понимает смещений: aware-дату он не принимает вовсе,
    а человек всё равно читает «19:00», а не «16:00 UTC»."""
    await held_screening(session)

    page = read(await exports.export(session, "screenings", MOSCOW))["Показы"]
    header = [cell.value for cell in page[1]]
    starts = [
        row[header.index("Начало")] for row in page.iter_rows(min_row=2, values_only=True)
    ]

    assert starts
    for value in starts:
        assert isinstance(value, datetime)
        assert value.tzinfo is None


async def test_screening_sheet_shows_the_attendance_share(session):
    _, _, screening, _, _ = await held_screening(session)

    page = read(await exports.export(session, "screenings", MOSCOW))["Показы"]
    header = [cell.value for cell in page[1]]
    row = next(
        row for row in page.iter_rows(min_row=2, values_only=True) if row[0] == screening.id
    )

    assert row[header.index("Записалось")] == 2
    assert row[header.index("Пришло")] == 1
    assert row[header.index("Дошло, %")] == 50


async def test_sheet_names_fit_excel_limits():
    taken: set[str] = set()
    long = sheet_name("Очень длинное имя листа, какого Excel не примет", taken)
    assert len(long) <= 31

    twin = sheet_name("Очень длинное имя листа, какого Excel не примет", taken)
    assert twin != long and len(twin) <= 31
    assert sheet_name("Пары: 1/8 [плей-офф]", set()) == "Пары  1 8  плей-офф"


def test_lists_become_readable_text():
    """Жанры и режиссёры лежат массивом — в ячейку он должен попасть строкой,
    а не питоновским `['Драма']`."""
    data = workbook([Sheet("Тест", ("Жанры",), [(["Драма", "Комедия"],)])], MOSCOW)
    assert read(data)["Тест"].cell(row=2, column=1).value == "Драма, Комедия"


def test_filename_says_what_and_when():
    name = exports.filename("users", datetime(2026, 9, 13))
    assert name.endswith(".xlsx") and "2026-09-13" in name
    assert "всё" in exports.filename(exports.EVERYTHING, datetime(2026, 9, 13))


# --- Пароль на выгрузку -----------------------------------------------------


def test_password_is_stored_hashed_and_checks_out():
    stored = analytics_gate.hash_password("очень секретно")

    assert "очень секретно" not in stored
    assert analytics_gate.verify(stored, "очень секретно")
    assert not analytics_gate.verify(stored, "очень секретн")


def test_two_identical_passwords_have_different_hashes():
    """Соль у каждого своя: иначе по одинаковым строкам в дампе видно,
    что пароль не меняли."""
    assert analytics_gate.hash_password("одинаковый") != analytics_gate.hash_password("одинаковый")


def test_no_password_means_closed_not_open():
    assert not analytics_gate.verify(None, "что угодно")
    assert not analytics_gate.verify("", "")
    assert not analytics_gate.verify("мусор-не-хеш", "мусор-не-хеш")


def test_short_password_is_refused():
    with pytest.raises(analytics_gate.GateError):
        analytics_gate.hash_password("1234")


async def test_password_can_be_set_and_removed(session):
    await analytics_gate.set_password(session, "выгрузка-2026", None)
    await session.commit()
    assert analytics_gate.verify(await analytics_gate.stored_hash(session), "выгрузка-2026")

    await analytics_gate.set_password(session, "", None)
    await session.commit()
    assert not await analytics_gate.is_set(session)


async def test_password_never_shows_up_in_settings(client, session):
    """Хеш лежит в той же таблице, что параметры §13, но вне реестра —
    и не должен просачиваться в форму параметров."""
    await analytics_gate.set_password(session, "выгрузка-2026", None)
    await session.commit()

    boss = await login(client, SUPERADMIN_TG_ID, "Босс")
    response = await client.get(
        "/api/admin/settings", headers={"Authorization": f"Bearer {boss['token']}"}
    )
    assert response.status_code == 200
    assert all(item["key"] != analytics_gate.SETTING_KEY for item in response.json())


def test_five_misses_close_the_door_for_a_while():
    gate = analytics_gate.AccessGate()
    now = datetime.now(UTC)

    for _ in range(analytics_gate.MAX_ATTEMPTS):
        assert gate.locked_for(1, now) is None
        gate.register_failure(1, now)

    assert gate.locked_for(1, now) is not None
    # Чужой перебор не запирает соседа.
    assert gate.locked_for(2, now) is None
    assert gate.locked_for(1, now + analytics_gate.LOCKOUT + timedelta(seconds=1)) is None


def test_unlock_expires_by_itself():
    gate = analytics_gate.AccessGate()
    now = datetime.now(UTC)

    gate.unlock(1, now)
    assert gate.is_unlocked(1, now + analytics_gate.UNLOCK_TTL - timedelta(minutes=1))
    assert not gate.is_unlocked(1, now + analytics_gate.UNLOCK_TTL)


def test_correct_password_forgives_earlier_misses():
    gate = analytics_gate.AccessGate()
    now = datetime.now(UTC)

    gate.register_failure(1, now)
    gate.register_failure(1, now)
    gate.unlock(1, now)

    assert gate.locked_for(1, now) is None


async def test_admin_sets_and_clears_the_password(client, session):
    boss = await login(client, SUPERADMIN_TG_ID, "Босс")
    headers = {"Authorization": f"Bearer {boss['token']}"}

    assert (await client.get("/api/admin/analytics/password", headers=headers)).json() == {
        "is_set": False,
        "updated_at": None,
        "updated_by": None,
    }

    saved = await client.put(
        "/api/admin/analytics/password", json={"password": "выгрузка-2026"}, headers=headers
    )
    assert saved.status_code == 200
    assert saved.json()["is_set"] is True
    assert saved.json()["updated_by"] == "Босс"
    assert analytics_gate.verify(await analytics_gate.stored_hash(session), "выгрузка-2026")

    cleared = await client.put(
        "/api/admin/analytics/password", json={"password": ""}, headers=headers
    )
    assert cleared.json()["is_set"] is False


async def test_short_password_is_refused_by_the_api(client):
    boss = await login(client, SUPERADMIN_TG_ID, "Босс")
    response = await client.put(
        "/api/admin/analytics/password",
        json={"password": "1234"},
        headers={"Authorization": f"Bearer {boss['token']}"},
    )
    assert response.status_code == 400


async def test_ordinary_user_cannot_touch_the_password(client):
    ordinary = await login(client, 777044, "Обычный")
    headers = {"Authorization": f"Bearer {ordinary['token']}"}

    assert (await client.get("/api/admin/analytics/password", headers=headers)).status_code == 403
    assert (
        await client.put(
            "/api/admin/analytics/password", json={"password": "подобранный"}, headers=headers
        )
    ).status_code == 403
